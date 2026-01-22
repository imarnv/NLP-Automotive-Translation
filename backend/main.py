from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, FileResponse
import uvicorn
import shutil
import os

app = FastAPI(title="Document Translation Backend")

# CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # Adjust in production
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Global progress state: {task_id: {"stage": "stage_name", "percent": 0}}
# For this simple single-user app, we use a single global.
progress_state = {"status": "idle", "percent": 0}

@app.get("/status")
async def health_check():
    return {"status": "ok", "service": "Backend is running"}

@app.get("/progress")
async def get_progress():
    return progress_state

def update_progress(stage, percent):
    global progress_state
    progress_state["status"] = stage
    progress_state["percent"] = percent

@app.post("/translate")
async def translate_document(
    file: UploadFile = File(...),
    target_lang: str = Form("Tamil") # Default to Tamil
):
    global progress_state
    update_progress("Starting...", 0)
    
    # Save uploaded file temporarily
    upload_dir = "uploads"
    os.makedirs(upload_dir, exist_ok=True)
    file_path = f"{upload_dir}/{file.filename}"
    output_filename = f"translated_{file.filename}.pdf"
    output_path = f"{upload_dir}/{output_filename}"
    
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
            
        update_progress("Preprocessing...", 10)
            
        # 1. Pipeline Start
        content = ""
        filename = file.filename.lower()
        
        try:
            if filename.endswith(".docx"):
                import docx
                doc = docx.Document(file_path)
                full_text = []
                for para in doc.paragraphs:
                    full_text.append(para.text)
                content = "\n".join(full_text)
            else:
                # Default to text/plain
                with open(file_path, "r", encoding="utf-8") as f:
                    content = f.read()
                    
        except Exception as e:
             print(f"Error reading file: {e}")
             content = "Error reading file. Please ensure it is a valid text or docx file."

        # 2. Preprocessing
        from backend.pipeline.preprocessing import load_glossary, classify_terms, segment_text, protect_terms
        
        update_progress("Loading Glossary...", 20)
        full_glossary = load_glossary("english_tamil_hindi_glossary.json")
        
        # Pass target_lang to classify
        protected_glossary, preferred_glossary = classify_terms(full_glossary, target_lang)
        
        update_progress("Segmenting Text...", 30)
        sentences = segment_text(content)
        
        # 3. Translation Loop
        from backend.pipeline.translation import translate_sentences
        from backend.pipeline.postprocessing import restore_placeholders, apply_preferred_translations
        
        full_translated_text = []
        total_sents = len(sentences)
        
        update_progress("Translating...", 40)
        
        # Determine actual lang code for translation function
        # Our classify_terms used 'ta'/'hi' but translate_sentences expects 'Tamil'/'Hindi' string to decide model
        # or we can pass the raw target_lang string which is fine as valid input for our translation.py logic
        
        for i, sent in enumerate(sentences):
            # Update progress based on sentence count (40% to 80%)
            if total_sents > 0:
                current_progress = 40 + int((i / total_sents) * 40)
            else:
                current_progress = 80
                
            if i % 5 == 0: # Update every 5 sentences to reduce overhead
                 update_progress(f"Translating ({i+1}/{total_sents})", current_progress)

            # Mask ONLY protected terms
            protected_sent, term_map = protect_terms(sent, protected_glossary)
            
            # Translate
            translated_batch = translate_sentences([protected_sent], target_lang=target_lang)
            translated_sent = translated_batch[0]
            
            # Restore Protected Placeholders
            restored_sent = restore_placeholders(translated_sent, term_map)
            
            # Apply Preferred Translations
            final_sent = apply_preferred_translations(restored_sent, preferred_glossary, target_lang=target_lang)
            
            full_translated_text.append(final_sent)
        
        update_progress("Post-processing...", 80)
        result_text = "\n".join(full_translated_text)
        
        # 4. PDF Generation
        update_progress("Generating PDF...", 90)
        from backend.utils.pdf_gen import generate_pdf
        generate_pdf(result_text, output_path)
        
        update_progress("Completed", 100)
        
        return {
            "filename": file.filename,
            "status": "Translation Complete",
            "download_url": f"/download/{output_filename}"
        }
    except Exception as e:
        update_progress("Error", 0)
        print(f"Server Error: {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/download/{filename}")
async def download_file(filename: str):
    file_path = f"uploads/{filename}"
    if os.path.exists(file_path):
        return FileResponse(file_path, media_type='application/pdf', filename=filename)
    raise HTTPException(status_code=404, detail="File not found")

if __name__ == "__main__":
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
